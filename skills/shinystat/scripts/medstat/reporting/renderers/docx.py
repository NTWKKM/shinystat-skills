"""
src/medstat/reporting/renderers/docx.py: Native Word (.docx) document renderer.
Applies ICMJE 3-rule publication table borders via OpenXML and inserts figures at exact 6.5 in width.
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor

from medstat.reporting.ir import (
    CalloutBlock,
    FigureBlock,
    HeadingBlock,
    ParagraphBlock,
    ReportDocument,
    TableBlock,
)


def _apply_three_rule_table_borders(table) -> None:
    """
    Applies publication 3-rule borders (ICMJE/NEJM standards) using OpenXML:
    Top border: 1.5pt solid black
    Bottom border: 1.5pt solid black
    Inside horizontal borders: Disabled (w:insideH set to none).
    Header divider: Applied via cell bottom borders on row 0 (1.0pt solid black).
    Vertical / left / right borders: NONE.
    """
    tblPr = table._tbl.tblPr
    # Remove existing tblBorders if present
    existing_borders = tblPr.find(qn("w:tblBorders"))
    if existing_borders is not None:
        tblPr.remove(existing_borders)

    borders_xml = parse_xml(
        f"<w:tblBorders {nsdecls('w')}>\n"
        f'  <w:top w:val="single" w:sz="12" w:space="0" w:color="000000"/>\n'
        f'  <w:left w:val="none"/>\n'
        f'  <w:bottom w:val="single" w:sz="12" w:space="0" w:color="000000"/>\n'
        f'  <w:right w:val="none"/>\n'
        f'  <w:insideH w:val="none"/>\n'
        f'  <w:insideV w:val="none"/>\n'
        f"</w:tblBorders>"
    )
    tblPr.append(borders_xml)

    # Format header row: bold and bottom 1.0pt line
    if len(table.rows) > 0:
        header_tr = table.rows[0]._tr
        trPr = header_tr.get_or_add_trPr()
        trPr.append(parse_xml(f"<w:tblHeader {nsdecls('w')}/>"))
        for cell in table.rows[0].cells:
            tcPr = cell._tc.get_or_add_tcPr()
            tcBorders = parse_xml(
                f"<w:tcBorders {nsdecls('w')}>\n"
                f'  <w:bottom w:val="single" w:sz="8" w:space="0" w:color="000000"/>\n'
                f"</w:tcBorders>"
            )
            tcPr.append(tcBorders)
            for p in cell.paragraphs:
                for run in p.runs:
                    run.bold = True


def render_docx(
    doc: ReportDocument,
    out_path: str | Path,
) -> str:
    """
    Renders ReportDocument into a native .docx Word document.
    """
    out_p = Path(out_path).resolve()
    out_p.parent.mkdir(parents=True, exist_ok=True)

    document = Document()

    # 1. Page Margins (Standard 1.0 inch)
    for section in document.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)

    # 2. Document Title
    title_p = document.add_paragraph()
    title_run = title_p.add_run(doc.title)
    title_run.font.name = "Arial"
    title_run.font.size = Pt(20)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(15, 23, 42)
    title_p.paragraph_format.space_after = Pt(4)

    # Metadata
    meta_parts = []
    if doc.authors:
        meta_parts.append(f"Authors: {', '.join(doc.authors)}")
    if doc.date:
        meta_parts.append(f"Date: {doc.date}")
    if doc.institution:
        meta_parts.append(f"Affiliation: {doc.institution}")

    if meta_parts:
        meta_p = document.add_paragraph()
        meta_run = meta_p.add_run(" | ".join(meta_parts))
        meta_run.font.name = "Arial"
        meta_run.font.size = Pt(9.5)
        meta_run.font.color.rgb = RGBColor(100, 116, 139)
        meta_p.paragraph_format.space_after = Pt(16)

    # 3. Process Blocks
    for block in doc.blocks:
        if isinstance(block, HeadingBlock):
            level = max(1, min(4, block.level))
            h_p = document.add_heading(block.text, level=level)
            h_p.paragraph_format.space_before = Pt(14)
            h_p.paragraph_format.space_after = Pt(4)
            for run in h_p.runs:
                run.font.name = "Arial"
                run.font.color.rgb = RGBColor(30, 41, 59)
        elif isinstance(block, ParagraphBlock):
            p = document.add_paragraph()
            p_run = p.add_run(block.text)
            p_run.font.name = "Arial"
            p_run.font.size = Pt(10.5)
            p.paragraph_format.line_spacing = 1.25
            p.paragraph_format.space_after = Pt(8)
        elif isinstance(block, CalloutBlock):
            c_p = document.add_paragraph()
            c_run = c_p.add_run(f"Note: {block.text}")
            c_run.font.name = "Arial"
            c_run.font.size = Pt(10.0)
            c_run.font.italic = True
            c_p.paragraph_format.left_indent = Inches(0.25)
            c_p.paragraph_format.space_after = Pt(10)
        elif isinstance(block, TableBlock):
            # Caption
            cap_p = document.add_paragraph()
            cap_run = cap_p.add_run(block.caption)
            cap_run.font.name = "Arial"
            cap_run.font.size = Pt(10.5)
            cap_run.font.bold = True
            cap_p.paragraph_format.space_before = Pt(12)
            cap_p.paragraph_format.space_after = Pt(4)

            # Table
            df = block.df
            n_rows, n_cols = len(df) + 1, len(df.columns)
            table = document.add_table(rows=n_rows, cols=n_cols)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER

            # Header row
            for c_idx, col_name in enumerate(df.columns):
                cell = table.cell(0, c_idx)
                cell.text = str(col_name)
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
                    for run in p.runs:
                        run.font.name = "Arial"
                        run.font.size = Pt(9.0)
                        run.font.bold = True

            # Data rows
            for r_idx, (_, row) in enumerate(df.iterrows()):
                for c_idx, col_name in enumerate(df.columns):
                    cell = table.cell(r_idx + 1, c_idx)
                    val = str(row[col_name]) if pd.notna(row[col_name]) else ""
                    cell.text = val
                    for p in cell.paragraphs:
                        p.paragraph_format.space_before = Pt(2)
                        p.paragraph_format.space_after = Pt(2)
                        for run in p.runs:
                            run.font.name = "Arial"
                            run.font.size = Pt(9.0)

            _apply_three_rule_table_borders(table)

            # Footnote
            if block.footnote:
                fn_p = document.add_paragraph()
                fn_run = fn_p.add_run(block.footnote)
                fn_run.font.name = "Arial"
                fn_run.font.size = Pt(8.5)
                fn_run.font.italic = True
                fn_run.font.color.rgb = RGBColor(100, 116, 139)
                fn_p.paragraph_format.space_after = Pt(12)
        elif isinstance(block, FigureBlock):
            fig_path = Path(block.figure.png_path)
            if not fig_path.exists():
                raise FileNotFoundError(f"Figure file not found: {fig_path}")
            fig_p = document.add_paragraph()
            fig_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            fig_p.paragraph_format.space_before = Pt(12)
            fig_p.paragraph_format.space_after = Pt(4)
            # Fit standard 6.5 inch width
            document.add_picture(str(fig_path), width=Inches(6.5))

            # Caption
            cap_p = document.add_paragraph()
            cap_run = cap_p.add_run(block.figure.caption)
            cap_run.font.name = "Arial"
            cap_run.font.size = Pt(9.5)
            cap_p.paragraph_format.space_after = Pt(14)

    document.save(str(out_p))
    return str(out_p)
