"""
tests/unit/test_report_ir_and_renderers.py: Unit tests for Report IR and Multi-Format Renderers.
"""

# ruff: noqa: E402

from __future__ import annotations

import sys
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parent.parent.parent
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

import pandas as pd
import pytest
from docx import Document
from pptx import Presentation

from medstat.figures.roc import plot_roc_curve
from medstat.reporting.ir import (
    ReportDocument,
)
from medstat.reporting.renderers import (
    render_docx,
    render_html,
    render_markdown,
    render_pptx,
)


@pytest.fixture
def sample_report_doc(tmp_path: Path) -> ReportDocument:
    """Creates a sample ReportDocument with synthetic clinical findings and a figure."""
    fig_path = tmp_path / "test_roc.png"
    roc_data = {
        "fpr": [0.0, 0.05, 0.15, 0.30, 0.60, 1.0],
        "tpr": [0.0, 0.50, 0.75, 0.88, 0.96, 1.0],
        "thresholds": [10.0, 8.0, 5.0, 3.0, 1.0, 0.1],
        "auc": 0.865,
        "ci_lower": 0.812,
        "ci_upper": 0.918,
        "optimal_point": {"fpr": 0.15, "tpr": 0.75, "cutoff": 5.0},
        "model_name": "Biomarker X ROC",
    }
    fig_result = plot_roc_curve(
        roc_data=roc_data,
        title="Biomarker X ROC Analysis",
        out_path=fig_path,
    )

    summary_df = pd.DataFrame(
        {
            "Variable": ["Age", "Biomarker X", "ICU Admission"],
            "Mean / Count": ["58.4 (12.1)", "2.45 (1.10)", "48 (24.0%)"],
            "P-value": ["0.042", "< 0.001", "-"],
        }
    )

    doc = ReportDocument(
        title="Clinical Validation of Biomarker X for ICU Triage",
        authors=["Dr. Jane Doe, MD", "Dr. John Smith, PhD"],
        date="2026-10-09",
        institution="Department of Critical Care Medicine",
        results_dict={
            "auc": 0.865,  # numerical result
            "sample_size": 200,
            "icu_count": 48,
            "icu_pct": 24.0,
            "age_mean": 58.4,
            "p_val": 0.042,
        },
    )

    doc.add_heading("Executive Summary", level=1)
    doc.add_paragraph(
        "In this observational cohort of 200 patients, Biomarker X showed significant prognostic value."
    )
    doc.add_callout(
        "Observational study design requires cautious interpretation regarding unmeasured confounding.",
        level="warning",
    )

    doc.add_heading("Baseline Cohort Characteristics", level=2)
    doc.add_table(
        summary_df,
        caption="Table 1: Baseline Characteristics of Study Cohort",
        footnote="Values presented as Mean (SD) or Count (%).",
    )

    doc.add_heading("Diagnostic Performance", level=2)
    doc.add_figure(fig_result)

    doc.add_heading("Methods & Protocol", level=2)
    doc.add_paragraph(
        "Statistical analysis was conducted using Python medstat package version 1.0 with 95% confidence intervals."
    )

    return doc


def test_report_ir_text_extraction(sample_report_doc: ReportDocument):
    """Verifies complete text extraction across all block types."""
    all_text = sample_report_doc.extract_all_text()
    assert "Clinical Validation of Biomarker X" in all_text
    assert "Executive Summary" in all_text
    assert "Table 1: Baseline Characteristics" in all_text
    assert "Values presented as Mean (SD)" in all_text
    assert "AUC" in all_text or "ROC" in all_text


def test_render_markdown(sample_report_doc: ReportDocument, tmp_path: Path):
    """Verifies markdown rendering and asset directory copying."""
    md_file = tmp_path / "report.md"
    out_path = render_markdown(sample_report_doc, md_file, assets_dir_name="figures")

    assert Path(out_path).exists()
    content = Path(out_path).read_text(encoding="utf-8")
    assert "# Clinical Validation of Biomarker X" in content
    assert "Dr. Jane Doe" in content
    assert "> [!NOTE]" in content
    assert "(figures/test_roc.png)" in content
    assert (tmp_path / "figures" / "test_roc.png").exists()


def test_render_html(sample_report_doc: ReportDocument, tmp_path: Path):
    """Verifies self-contained HTML generation with inline base64 image."""
    html_file = tmp_path / "report.html"
    out_path = render_html(sample_report_doc, html_file)

    assert Path(out_path).exists()
    content = Path(out_path).read_text(encoding="utf-8")
    assert "<!DOCTYPE html>" in content
    assert "data:image/png;base64," in content
    assert "border-top: 2px solid #000000" in content
    assert "border-bottom: 2px solid #000000" in content
    assert "Clinical Validation of Biomarker X" in content


def test_render_docx(sample_report_doc: ReportDocument, tmp_path: Path):
    """Verifies native Word .docx rendering with 3-rule borders and images."""
    docx_file = tmp_path / "report.docx"
    out_path = render_docx(sample_report_doc, docx_file)

    assert Path(out_path).exists()
    assert Path(out_path).stat().st_size > 1000

    # Load with python-docx and inspect elements
    doc = Document(str(out_path))
    assert len(doc.paragraphs) > 5
    assert len(doc.tables) >= 1
    table = doc.tables[0]
    assert len(table.rows) == 4  # header + 3 data rows
    assert "Biomarker X" in table.rows[2].cells[0].text


def test_render_pptx(sample_report_doc: ReportDocument, tmp_path: Path):
    """Verifies native PowerPoint .pptx 16:9 widescreen rendering."""
    pptx_file = tmp_path / "presentation.pptx"
    out_path = render_pptx(sample_report_doc, pptx_file)

    assert Path(out_path).exists()
    assert Path(out_path).stat().st_size > 1000

    prs = Presentation(str(out_path))
    # Check 16:9 widescreen dimensions (13.333 in x 7.5 in)
    assert pytest.approx(prs.slide_width.inches, rel=1e-2) == 13.333
    assert pytest.approx(prs.slide_height.inches, rel=1e-2) == 7.5

    # Check slides: title slide, content slide, table slide, figure slide
    assert len(prs.slides) >= 4


def test_render_pdf(sample_report_doc: ReportDocument, tmp_path: Path):
    """Verifies publication PDF rendering via Playwright or soffice fallback."""
    import importlib.util
    import shutil

    has_playwright = importlib.util.find_spec("playwright") is not None
    has_soffice = bool(shutil.which("soffice"))

    if not has_playwright and not has_soffice:
        pytest.skip("Neither Playwright nor LibreOffice 'soffice' is available.")

    from medstat.reporting.renderers import render_pdf

    pdf_file = tmp_path / "report.pdf"
    try:
        out_path = render_pdf(sample_report_doc, pdf_file)
    except RuntimeError as e:
        if "Playwright" in str(e) or "soffice" in str(e):
            pytest.skip(f"PDF engine runtime unavailable: {e}")
        raise

    assert Path(out_path).exists()
    assert Path(out_path).stat().st_size > 5000
    assert Path(out_path).read_bytes().startswith(b"%PDF")


def test_renderers_missing_figure_raises_error(tmp_path: Path):
    """Verifies that missing figure file raises FileNotFoundError consistently across renderers."""
    from medstat.figures.base import FigureResult
    from medstat.reporting.renderers import (
        render_docx,
        render_html,
        render_markdown,
        render_pptx,
    )

    missing_fig = FigureResult(
        png_path=str(tmp_path / "non_existent_figure.png"),
        alt_text="Missing figure",
        caption="Missing figure caption",
        source_df=pd.DataFrame({"x": [1]}),
        csv_path=str(tmp_path / "dummy.csv"),
    )
    doc = ReportDocument(title="Test Missing Figure")
    doc.add_figure(missing_fig)

    with pytest.raises(FileNotFoundError):
        render_markdown(doc, tmp_path / "test.md")

    with pytest.raises(FileNotFoundError):
        render_html(doc, tmp_path / "test.html")

    with pytest.raises(FileNotFoundError):
        render_docx(doc, tmp_path / "test.docx")

    with pytest.raises(FileNotFoundError):
        render_pptx(doc, tmp_path / "test.pptx")


def test_render_pptx_table_pagination(tmp_path: Path):
    """Verifies that tables with more than 12 rows paginate across multiple slides."""
    df_large = pd.DataFrame(
        {
            "Variable": [f"Covariate_{i + 1:02d}" for i in range(25)],
            "Value": [round(10.0 + i * 1.5, 2) for i in range(25)],
        }
    )
    doc = ReportDocument(title="Pagination Test Report")
    doc.add_heading("Baseline Patient Demographics", level=1)
    doc.add_table(df_large, caption="Baseline Patient Demographics")
    pptx_file = tmp_path / "paginated_presentation.pptx"
    out_path = render_pptx(doc, pptx_file)

    assert Path(out_path).exists()
    prs = Presentation(str(out_path))

    # Expect: Title slide + 3 continuation table slides (12 + 12 + 1 rows)
    assert len(prs.slides) == 4

    # Verify slide titles and footnotes
    expected_titles = [
        "Baseline Patient Demographics (Part 1/3)",
        "Baseline Patient Demographics (Part 2/3)",
        "Baseline Patient Demographics (Part 3/3)",
    ]
    expected_notes = [
        "Rows 1–12 of 25",
        "Rows 13–24 of 25",
        "Rows 25–25 of 25",
    ]

    for idx, slide in enumerate(list(prs.slides)[1:]):
        # Header title
        header_shapes = [
            s
            for s in slide.shapes
            if s.has_text_frame and len(s.text_frame.paragraphs) >= 2
        ]
        assert len(header_shapes) >= 1
        assert header_shapes[0].text_frame.paragraphs[1].text == expected_titles[idx]

        # Footnote
        note_shapes = [
            s
            for s in slide.shapes
            if s.has_text_frame and expected_notes[idx] in s.text_frame.text
        ]
        assert len(note_shapes) >= 1

        # Table rows: header + chunk data rows
        table_shapes = [s for s in slide.shapes if s.has_table]
        assert len(table_shapes) == 1
        tbl = table_shapes[0].table
        if idx == 0 or idx == 1:
            assert len(tbl.rows) == 13  # 1 header + 12 data rows
        else:
            assert len(tbl.rows) == 2  # 1 header + 1 data row
